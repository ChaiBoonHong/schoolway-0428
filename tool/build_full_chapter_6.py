import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.bold = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Times New Roman'
        r_bold.font.size = Pt(12)
        r_bold.font.bold = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.italic = italic
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.bold = True
    return p

def add_image_centered(doc, img_path, width_in_inches=5.8):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img_path, width=Inches(width_in_inches))
    else:
        print(f"Warning: Image path does not exist: {img_path}")

def add_test_case_table(doc, use_case, func_id, date_created, role, steps, test_desc=None, trials_desc=None, fig_ref=None):
    if test_desc:
        add_body_paragraph(doc, test_desc)
    
    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="CCCCCC")
    
    meta_data = [
        ("Use Case", use_case),
        ("Function ID", func_id),
        ("Date Created", date_created),
        ("Role", role)
    ]
    
    for idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_k = row.cells[0]
        cell_v = row.cells[1]
        
        cell_k.width = Inches(2.0)
        cell_v.width = Inches(4.5)
        
        set_cell_background(cell_k, "F2F4F7")
        set_cell_margins(cell_k, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_v, top=60, bottom=60, left=100, right=100)
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_before = Pt(0)
        pk.paragraph_format.space_after = Pt(0)
        rk = pk.add_run(k)
        rk.font.name = 'Times New Roman'
        rk.font.size = Pt(10.5)
        rk.font.bold = True
        
        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_before = Pt(0)
        pv.paragraph_format.space_after = Pt(0)
        rv = pv.add_run(v)
        rv.font.name = 'Times New Roman'
        rv.font.size = Pt(10.5)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    
    # Steps Table
    steps_table = doc.add_table(rows=len(steps) + 1, cols=5)
    steps_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(steps_table, color="CCCCCC")
    
    headers = ["No.", "Input Values", "Expected Results", "Actual Results", "Pass / Fail"]
    col_widths = [Inches(0.5), Inches(1.8), Inches(1.9), Inches(1.6), Inches(0.7)]
    
    hdr_row = steps_table.rows[0]
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_background(cell, "02569B")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if c_idx in [0, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, step in enumerate(steps):
        row = steps_table.rows[r_idx + 1]
        bg_color = "FAFBFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(step):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if c_idx in [0, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            if c_idx == 4:
                r.font.bold = True
                if "Pass" in str(val):
                    r.font.color.rgb = RGBColor(42, 123, 80)
                else:
                    r.font.color.rgb = RGBColor(192, 57, 43)

    if trials_desc:
        p_res = doc.add_paragraph()
        p_res.paragraph_format.space_before = Pt(4)
        p_res.paragraph_format.space_after = Pt(6)
        r_res = p_res.add_run(trials_desc)
        r_res.font.name = 'Times New Roman'
        r_res.font.size = Pt(11)
        r_res.font.italic = True
        
    if fig_ref:
        p_fig = doc.add_paragraph()
        p_fig.paragraph_format.space_before = Pt(0)
        p_fig.paragraph_format.space_after = Pt(10)
        r_fig = p_fig.add_run(f"Verification Proof: As evidenced in {fig_ref}.")
        r_fig.font.name = 'Times New Roman'
        r_fig.font.size = Pt(11)
        r_fig.font.bold = True
        r_fig.font.color.rgb = RGBColor(2, 86, 155)

print("Table builder functions ready.")
